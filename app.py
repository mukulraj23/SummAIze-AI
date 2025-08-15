import os
import math
import io # Used for in-memory file handling

# New imports for file handling and PDF generation
import pandas as pd
import docx
from fpdf import FPDF

import google.generativeai as genai
from flask import Flask, render_template, request, jsonify, send_file
from dotenv import load_dotenv
import PyPDF2

# Load environment variables
load_dotenv()

# Configure Google AI
try:
    genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))
except Exception as e:
    print(f"Error configuring Google AI: {e}")

model = genai.GenerativeModel('gemini-1.5-flash-latest')
app = Flask(__name__)

# --- Text Extraction Functions ---
def extract_text_from_pdf(file_stream):
    pdf_reader = PyPDF2.PdfReader(file_stream)
    text = ""
    num_pages = len(pdf_reader.pages)
    for page in pdf_reader.pages:
        text += page.extract_text() or ""
    return text, num_pages

def extract_text_from_docx(file_stream):
    doc = docx.Document(file_stream)
    text = "\n".join([para.text for para in doc.paragraphs])
    word_count = len(text.split())
    num_pages = math.ceil(word_count / 500) # Estimate pages
    return text, num_pages

def extract_text_from_excel(file_stream):
    df = pd.read_excel(file_stream, engine='openpyxl', sheet_name=None)
    text = ""
    for sheet_name, sheet_df in df.items():
        text += f"--- Sheet: {sheet_name} ---\n"
        text += sheet_df.to_string(index=False)
        text += "\n\n"
    word_count = len(text.split())
    num_pages = math.ceil(word_count / 500) # Estimate pages
    return text, num_pages

# --- Flask Routes ---
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/summarize', methods=['POST'])
def summarize():
    text_to_summarize = ""
    num_pages = 0

    if 'document_input' in request.files and request.files['document_input'].filename != '':
        file = request.files['document_input']
        filename = file.filename.lower()

        try:
            if filename.endswith('.pdf'):
                text_to_summarize, num_pages = extract_text_from_pdf(file)
            elif filename.endswith('.docx'):
                text_to_summarize, num_pages = extract_text_from_docx(file)
            elif filename.endswith(('.xlsx', '.xls')):
                text_to_summarize, num_pages = extract_text_from_excel(file)
            elif filename.endswith('.txt'):
                text_to_summarize = file.read().decode('utf-8')
                word_count = len(text_to_summarize.split())
                num_pages = math.ceil(word_count / 500)
            else:
                return jsonify({'error': 'Unsupported file format. Please use .txt, .pdf, .docx, or .xlsx.'}), 400
        except Exception as e:
            return jsonify({'error': f'Error processing file: {str(e)}'}), 500

    elif 'text_input' in request.form and request.form['text_input']:
        text_to_summarize = request.form['text_input']
        word_count = len(text_to_summarize.split())
        num_pages = math.ceil(word_count / 500)

    if not text_to_summarize.strip():
        return jsonify({'error': 'No text or valid document provided.'}), 400

    if num_pages <= 1: points = "3 to 4"
    elif num_pages <= 2: points = "6 to 7"
    else: points = f"{int(num_pages * 3)} to {int(num_pages * 3 + 2)}"

    prompt = f"Analyze the document and summarize its key information into a numbered list of {points} important points. Each point must be clear and concise. Provide only the numbered list.\n\nDOCUMENT:\n---\n{text_to_summarize}"

    try:
        response = model.generate_content(prompt)
        summary = response.text
        return jsonify({'summary': summary})
    except Exception as e:
        return jsonify({'error': f'An error occurred with the AI model: {str(e)}'}), 500

# NEW: Route for downloading the summary as a PDF
@app.route('/download_pdf', methods=['POST'])
def download_pdf():
    summary_text = request.json.get('summary_text', '')
    if not summary_text:
        return jsonify({'error': 'No summary text provided'}), 400

    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        # Add a title
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, 'AI Generated Summary', 0, 1, 'C')
        pdf.ln(10) # Add a little space
        
        # Add the summary text
        pdf.set_font("Arial", size=12)
        # multi_cell is important for handling long text and line breaks
        # We need to encode the text properly for FPDF
        encoded_text = summary_text.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 10, encoded_text)

        # Create an in-memory buffer for the PDF
        buffer = io.BytesIO()
        pdf.output(buffer)
        buffer.seek(0)

        return send_file(
            buffer,
            as_attachment=True,
            download_name='summary.pdf',
            mimetype='application/pdf'
        )
    except Exception as e:
        return jsonify({'error': f'Failed to generate PDF: {str(e)}'}), 500

if __name__ == '__main__':
    app.run(debug=True)